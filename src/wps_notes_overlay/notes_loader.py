"""Load per-slide speaker notes from OOXML presentation files.

For large decks (hundreds of MB of media), prefer ``load_deck_light`` which
only reads slide relationship + notesSlide XML from the ZIP — it never
inflates images or builds a full python-pptx object graph.

WPS on Mac usually saves as .pptx (OOXML). Older .dps is proprietary and
often not readable as ZIP/OOXML — we detect that and ask the user to
「另存为 .pptx」.
"""

from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".pptx", ".ppsx", ".pptm", ".potx"}

# OOXML namespaces commonly seen in pptx
_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}
_NOTES_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
)
_SLIDE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
)


@dataclass(frozen=True)
class SlideNotes:
    index: int  # 1-based
    notes: str
    title_hint: str = ""


@dataclass
class DeckNotes:
    path: Path
    slides: list[SlideNotes]

    @property
    def count(self) -> int:
        return len(self.slides)

    def get(self, index_1based: int) -> SlideNotes | None:
        if 1 <= index_1based <= len(self.slides):
            return self.slides[index_1based - 1]
        return None


def looks_like_ooxml_zip(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
        return "[Content_Types].xml" in names and any(
            n.startswith("ppt/") for n in names
        )
    except (zipfile.BadZipFile, OSError):
        return False


def _local(tag: str) -> str:
    if "}" in tag:
        return tag.rsplit("}", 1)[-1]
    return tag


def _text_from_notes_xml(xml_bytes: bytes) -> str:
    """Extract speaker-notes body text from a notesSlide part.

    Prefer shapes whose placeholder type is ``body``; fall back to all ``a:t``
    nodes excluding a lone numeric slide-number shape when possible.
    """
    root = ET.fromstring(xml_bytes)
    body_chunks: list[str] = []
    other_chunks: list[str] = []

    for sp in root.iter():
        if _local(sp.tag) != "sp":
            continue
        ph_type = None
        for el in sp.iter():
            if _local(el.tag) == "ph":
                ph_type = el.attrib.get("type") or el.attrib.get(
                    "{http://schemas.openxmlformats.org/presentationml/2006/main}type"
                )
                # sometimes type is bare
                if ph_type is None:
                    ph_type = el.attrib.get("type", "")
                break
        texts: list[str] = []
        for el in sp.iter():
            if _local(el.tag) == "t" and el.text:
                texts.append(el.text)
            if _local(el.tag) == "t" and el.tail and el.tail.strip():
                pass
        # also join with breaks: collect paragraph-wise
        paras: list[str] = []
        for el in sp.iter():
            if _local(el.tag) == "p":
                parts: list[str] = []
                for t in el.iter():
                    if _local(t.tag) == "t" and t.text:
                        parts.append(t.text)
                if parts:
                    paras.append("".join(parts))
        block = "\n".join(paras).strip() if paras else "".join(texts).strip()
        if not block:
            continue
        if ph_type == "body":
            body_chunks.append(block)
        elif ph_type in ("sldNum", "slideNumber", "hdr", "ftr", "title"):
            continue
        else:
            other_chunks.append(block)

    if body_chunks:
        return "\n\n".join(body_chunks).strip()
    # Fallback: all non-trivial text that isn't just a slide number
    kept = [c for c in other_chunks if not re.fullmatch(r"\d+", c.strip())]
    return "\n\n".join(kept).strip()


def _looks_like_boilerplate(text: str) -> bool:
    lower = text.lower()
    if "all rights reserved" in lower or "©" in text or "(c)" in lower:
        return True
    if re.fullmatch(r"\d{1,3}", text.strip()):
        return True
    return False


def _title_hint_from_slide_xml(xml_bytes: bytes) -> str:
    root = ET.fromstring(xml_bytes)
    candidates: list[str] = []
    for sp in root.iter():
        if _local(sp.tag) != "sp":
            continue
        ph_type = None
        for el in sp.iter():
            if _local(el.tag) == "ph":
                ph_type = el.attrib.get("type", "")
                break
        parts: list[str] = []
        for el in sp.iter():
            if _local(el.tag) == "t" and el.text:
                parts.append(el.text)
        text = "".join(parts).strip()
        if not text:
            continue
        line = text.splitlines()[0].strip()[:80]
        if ph_type in ("title", "ctrTitle") and not _looks_like_boilerplate(line):
            return line
        if not _looks_like_boilerplate(line):
            candidates.append(line)
    return candidates[0] if candidates else ""


def _resolve_rel_target(base_dir: str, target: str) -> str:
    # base_dir like ppt/slides ; target like ../notesSlides/notesSlide1.xml
    parts = base_dir.split("/") + target.split("/")
    stack: list[str] = []
    for part in parts:
        if part in ("", "."):
            continue
        if part == "..":
            if stack:
                stack.pop()
            continue
        stack.append(part)
    return "/".join(stack)


def _slide_order(zf: zipfile.ZipFile) -> list[str]:
    """Return package paths of slides in presentation order."""
    try:
        pres = zf.read("ppt/presentation.xml")
    except KeyError as exc:
        raise ValueError("缺少 ppt/presentation.xml，不是有效的演示文稿包") from exc
    root = ET.fromstring(pres)
    r_id_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    rids: list[str] = []
    for el in root.iter():
        if _local(el.tag) != "sldId":
            continue
        rid = el.attrib.get(r_id_attr)
        if not rid:
            for k, v in el.attrib.items():
                if "relationships" in k and _local(k) == "id":
                    rid = v
                    break
        if rid:
            rids.append(rid)

    try:
        rels_xml = zf.read("ppt/_rels/presentation.xml.rels")
    except KeyError as exc:
        raise ValueError("缺少 ppt/_rels/presentation.xml.rels") from exc
    rel_root = ET.fromstring(rels_xml)
    rid_to_target: dict[str, str] = {}
    for rel in rel_root:
        if _local(rel.tag) != "Relationship":
            continue
        rid_to_target[rel.attrib.get("Id", "")] = rel.attrib.get("Target", "")

    slides: list[str] = []
    for rid in rids:
        target = rid_to_target.get(rid)
        if not target:
            continue
        path = _resolve_rel_target("ppt", target)
        norm = path.replace("\\", "/")
        if "/slides/slide" in norm and norm.endswith(".xml"):
            slides.append(path)
    if slides:
        return slides

    names = [
        n
        for n in zf.namelist()
        if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)
    ]

    def key(n: str) -> int:
        m = re.search(r"slide(\d+)\.xml$", n)
        return int(m.group(1)) if m else 0

    return sorted(names, key=key)


def load_deck_light(path: str | Path) -> DeckNotes:
    """Parse notes using ZIP + XML only (media streams never decoded)."""
    path = Path(path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"找不到文件: {path}")

    suffix = path.suffix.lower()
    if suffix == ".dps" and not looks_like_ooxml_zip(path):
        raise ValueError(
            "检测到旧版 WPS .dps（专有格式），本工具无法直接读取备注。\n"
            "请在 WPS 中打开后「另存为」→ 选择 PowerPoint 演示文稿 (*.pptx)，再加载该 .pptx。"
        )
    if suffix not in SUPPORTED_SUFFIXES and suffix != ".dps" and not looks_like_ooxml_zip(path):
        raise ValueError(
            f"不支持的格式: {suffix or '(无扩展名)'}。请使用 .pptx（推荐）或让 WPS 另存为 .pptx。"
        )

    try:
        zf_ctx = zipfile.ZipFile(path)
    except zipfile.BadZipFile as exc:
        raise ValueError(f"无法作为 OOXML(ZIP) 打开: {path}") from exc

    with zf_ctx as zf:
        if "[Content_Types].xml" not in zf.namelist():
            raise ValueError(f"不是有效的 OOXML 包: {path}")
        slide_paths = _slide_order(zf)
        if not slide_paths:
            raise ValueError("演示文稿里没有幻灯片。")

        slides: list[SlideNotes] = []
        for i, slide_path in enumerate(slide_paths, start=1):
            title_hint = ""
            try:
                title_hint = _title_hint_from_slide_xml(zf.read(slide_path))
            except KeyError:
                pass
            except ET.ParseError:
                pass

            notes = ""
            # slide rels path
            # ppt/slides/slide1.xml -> ppt/slides/_rels/slide1.xml.rels
            parent, name = slide_path.rsplit("/", 1)
            rels_path = f"{parent}/_rels/{name}.rels"
            try:
                rels_xml = zf.read(rels_path)
                rel_root = ET.fromstring(rels_xml)
                notes_target = None
                for rel in rel_root:
                    if _local(rel.tag) != "Relationship":
                        continue
                    if rel.attrib.get("Type") == _NOTES_REL_TYPE:
                        notes_target = rel.attrib.get("Target")
                        break
                if notes_target:
                    notes_path = _resolve_rel_target(parent, notes_target)
                    try:
                        notes = _text_from_notes_xml(zf.read(notes_path))
                    except KeyError:
                        notes = ""
                    except ET.ParseError:
                        notes = ""
            except KeyError:
                notes = ""
            except ET.ParseError:
                notes = ""

            slides.append(SlideNotes(index=i, notes=notes, title_hint=title_hint))

    return DeckNotes(path=path, slides=slides)


def load_deck(path: str | Path, *, prefer_light: bool = True) -> DeckNotes:
    """Load deck notes. Default: light ZIP parser (safe for huge media-heavy pptx)."""
    path = Path(path).expanduser().resolve()
    if prefer_light:
        return load_deck_light(path)

    # Fallback heavy path (python-pptx) — only for odd packages light parser rejects
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    try:
        prs = Presentation(str(path))
    except PackageNotFoundError as exc:
        raise ValueError(f"无法作为 OOXML 演示文稿打开: {path}") from exc

    slides: list[SlideNotes] = []
    for i, slide in enumerate(prs.slides, start=1):
        title = ""
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = (shape.text_frame.text or "").strip()
            if text:
                title = text.splitlines()[0].strip()[:80]
                break
        notes = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                notes = (frame.text or "").strip()
        slides.append(SlideNotes(index=i, notes=notes, title_hint=title))
    if not slides:
        raise ValueError("演示文稿里没有幻灯片。")
    return DeckNotes(path=path, slides=slides)
